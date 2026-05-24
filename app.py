import sys
import html
import base64
import io
import json
import os
from typing import Any
from dataclasses import dataclass, field

try:
    import streamlit as st
    STREAMLIT_AVAILABLE = True
except ImportError:
    STREAMLIT_AVAILABLE = False
    print("CRITICAL: Streamlit is required. Run: pip install streamlit")
    sys.exit(1)

try:
    from parser import extract_text
    from analyser import analyse_resume, generate_cv, generate_cover_letter
    BACKEND_AVAILABLE = True
except Exception as e:
    print(f"\n>>> CRITICAL BACKEND ERROR: {e} <<<\n")
    BACKEND_AVAILABLE = False

_PAGE_CONFIG: dict[str, str] = {
    "page_title":            "RAGsToRiches",
    "page_icon":             "RT",
    "layout":                "wide",
    "initial_sidebar_state": "collapsed",
}

_SCORE_THRESHOLDS: dict[str, dict[str, Any]] = {
    "strong": {"min": 70, "color": "#15C39A", "label": "Strong"},
    "medium": {"min": 50, "color": "#E8A735", "label": "Needs work"},
    "weak":   {"min": 0,  "color": "#E5534B", "label": "Needs improvement"},
}

_CUSTOM_CSS: str = """
@import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;700&display=swap');
@import url('https://fonts.googleapis.com/css2?family=Instrument+Serif:ital@0;1&display=swap');

/* Base */
*, *::before, *::after { box-sizing: border-box; }
#MainMenu, footer, header { visibility: hidden; }
html, body, [class*="css"] { font-family: 'DM Sans', sans-serif; color: #0D0F11; background: #F7F8FA; }
.main .block-container { padding: 1.5rem 2.5rem 3rem; max-width: 1440px; margin: 0 auto; }

/* Header */
.hero-wrap {
    background: #0a0a0f;
    border-radius: 0 0 28px 28px;
    padding: 5rem 3rem 5rem;
    margin: -1rem -3rem 2.5rem;
    position: relative;
    overflow: hidden;
}
.hero-wrap::before {
    content: '';
    position: absolute;
    top: -120px; right: -140px;
    width: 700px; height: 700px;
    background: radial-gradient(circle, rgba(139,92,246,0.18) 0%, transparent 70%);
    filter: blur(60px);
    pointer-events: none;
}
.hero-wrap::after {
    content: '';
    position: absolute;
    bottom: -120px; left: -40px;
    width: 700px; height: 700px;
    background: radial-gradient(circle, rgba(20,184,166,0.12) 0%, transparent 70%);
    filter: blur(60px);
    pointer-events: none;
}
.hero-eyebrow {
    font-size: 0.72rem;
    font-weight: 600;
    letter-spacing: 0.16em;
    text-transform: uppercase;
    color: #14b8a6;
    margin-bottom: 0.75rem;
}
.hero-title {
    color: #f8fafc !important;
    font-family: 'Instrument Serif', serif !important;
    font-weight: 400 !important;
    font-size: 72px !important;
    font-style: normal !important;
    line-height: 1 !important;
    letter-spacing: -0.02em !important;
    -webkit-font-smoothing: antialiased !important;
    text-rendering: optimizeLegibility !important;
}

.hero-title .title-prefix {
    font-style: normal;
    color: #f8fafc;
}
.hero-title .accent {
    font-style: italic;
    color: #a78bfa;
}
.hero-sub {
    font-size: 1.08rem;
    color: rgba(255,255,255,0.72);
    font-weight: 400;
    margin: 0;
    max-width: 640px;
}
.hero-badge {
    display: inline-flex;
    align-items: center;
    gap: 6px;
    background: rgba(20,184,166,0.12);
    border: 1px solid rgba(20,184,166,0.25);
    color: #14b8a6;
    border-radius: 999px;
    padding: 0.3rem 0.85rem;
    font-size: 0.75rem;
    font-weight: 500;
    margin-top: 1.25rem;
}

/* ── UI Elements & Labels ── */
.section-label {
    font-family: 'Instrument Serif', serif !important;
    font-weight: 400 !important;
    font-size: 0.95rem !important;
    font-style: normal !important;
    text-transform: uppercase;
    letter-spacing: 0.02em !important;
    color: #6B7280; margin-bottom: 0.5rem; display: block;
    -webkit-font-smoothing: antialiased !important;
    text-rendering: optimizeLegibility !important;
}
.slim-divider {
    border: 0;
    height: 1px;
    margin: 0.9rem 0 1.15rem;
    background: linear-gradient(90deg, rgba(255,255,255,0), rgba(255,255,255,0.75), rgba(255,255,255,0));
    box-shadow: 0 1px 0 rgba(255,255,255,0.08);
}
.result-section-head {
    font-family: 'Instrument Serif', serif !important;
    font-weight: 400 !important;
    font-size: 0.95rem !important;
    font-style: normal !important;
    color: #0D0F11;
    padding: 0.75rem 0 0.4rem; display: flex; align-items: center; gap: 0.55rem;
    letter-spacing: 0.02em !important;
    -webkit-font-smoothing: antialiased !important;
    text-rendering: optimizeLegibility !important;
}
.result-section-head::after { content: ''; flex: 1; height: 1px; background: #E8EAED; }

/* Cards */
.card, .score-card {
    background: #FFFFFF; border: 1px solid #E8EAED; border-radius: 16px;
    padding: 1.25rem; margin-bottom: 1rem; box-shadow: 0 1px 3px rgba(13,15,17,0.04);
    transition: box-shadow 0.2s ease;
}
.card:hover { box-shadow: 0 4px 12px rgba(13,15,17,0.07); }

/* Setup Band */
.setup-band {
    border: 0;
    background: transparent;
    padding: 0;
    margin: 0 0 1rem;
    box-shadow: none;
}

/* Score */
.score-ring-wrap { display: flex; align-items: center; gap: 1rem; padding: 0.25rem 0; }
.score-number {
    font-family: 'Instrument Serif', serif !important;
    font-size: 4rem;
    line-height: 1;
    font-weight: 400;
    letter-spacing: -0.03em;
}
.score-label-text {
    font-family: 'Instrument Serif', serif !important;
    font-weight: 400 !important;
    font-size: 0.95rem !important;
    font-style: normal !important;
    color: #0D0F11;
    letter-spacing: 0.02em !important;
    -webkit-font-smoothing: antialiased !important;
    text-rendering: optimizeLegibility !important;
}
.score-sub, .muted { font-size: 0.8rem; color: #6B7280; }

/* Contact & Status */
.contact-grid, .kw-wrap, .review-toolbar { display: flex; flex-wrap: wrap; gap: 0.5rem; align-items: center; }
.contact-chip, .status-chip {
    background: #F7F8FA; border: 1px solid #E8EAED; border-radius: 999px;
    padding: 0.28rem 0.7rem; font-size: 0.76rem; color: #0D0F11;
    transition: background 0.15s;
}
.contact-chip:hover, .status-chip:hover { background: #E8EAED; }
.contact-chip b { color: #6B7280; font-weight: 700; margin-right: 0.25rem; }

/* PDF Viewer */
.pdf-shell {
    background: #2D3136; border: 1px solid #3C4148; border-radius: 14px;
    overflow: hidden; min-height: 720px; box-shadow: 0 4px 20px rgba(13,15,17,0.15);
}
.pdf-frame { width: 100%; min-height: 720px; border: 0; display: block; background: #2D3136; }

/* Suggestion Cards */
.suggestion-scroll { max-height: 760px; overflow-y: auto; padding-right: 0.3rem; }
.suggestion-scroll::-webkit-scrollbar { width: 5px; }
.suggestion-scroll::-webkit-scrollbar-thumb { background: #cbd5e1; border-radius: 999px; }
.suggestion-card {
    background: #FFFFFF; border: 1px solid #E8EAED; border-radius: 12px;
    padding: 0; overflow: hidden; margin-bottom: 0.85rem;
    box-shadow: 0 1px 3px rgba(13,15,17,0.04); transition: all 0.2s ease;
}
.suggestion-card:hover { box-shadow: 0 4px 16px rgba(13,15,17,0.08); transform: translateY(-1px); }
.suggestion-card.accepted { border-color: #15C39A; background: #E6F8F3; }
.suggestion-card.dismissed { border-color: #fca5a5; background: #fef2f2; }
.suggestion-head { padding: 0.8rem 1rem 0; display: flex; justify-content: space-between; gap: 0.75rem; align-items: center; }
.suggestion-title { font-size: 0.72rem; font-weight: 700; color: #6B7280; text-transform: uppercase; letter-spacing: 0.08em; }
.fw-badge {
    background: rgba(21,195,154,0.1); color: #15C39A;
    font-size: 0.67rem; font-weight: 700; padding: 0.2rem 0.55rem;
    border-radius: 999px; white-space: nowrap; border: 1px solid rgba(21,195,154,0.3);
}
.severity-dot {
    width: 8px; height: 8px; border-radius: 50%; display: inline-block;
    margin-right: 0.35rem; vertical-align: middle; flex-shrink: 0;
}
.severity-dot.red { background: #E5534B; }
.severity-dot.yellow { background: #E8A735; }
.severity-dot.green { background: #15C39A; }

/* Rewrite Grid */
.rewrite-grid {
    display: grid; grid-template-columns: minmax(0, 1fr) minmax(0, 1fr);
    gap: 0; border-top: 1px solid #E8EAED; margin-top: 0.7rem;
}
.rewrite-pane { padding: 0.85rem 1rem; min-width: 0; }
.rewrite-pane.before { background: #FAFAFA; border-right: 1px solid #E8EAED; }
.rewrite-pane.after { background: #FFFFFF; }
.pane-label {
    display: block; font-size: 0.64rem; font-weight: 700;
    letter-spacing: 0.1em; text-transform: uppercase; color: #6B7280; margin-bottom: 0.35rem;
}
.rewrite-text { font-size: 0.84rem; line-height: 1.6; color: #0D0F11; margin: 0; overflow-wrap: anywhere; }
.reasoning-row {
    padding: 0.72rem 1rem; border-top: 1px solid #E8EAED;
    background: #F7F8FA; color: #6B7280; font-size: 0.78rem; line-height: 1.5;
    font-style: italic;
}
.decision-bar { padding: 0.5rem 1rem; border-top: 1px solid #E8EAED; font-size: 0.74rem; font-weight: 700; }
.decision-bar.accepted { background: #15C39A; color: #FFFFFF; }
.decision-bar.dismissed { background: #fee2e2; color: #991b1b; }

/* Spacing between decision-bar/card and accept-dismiss buttons */
.suggestion-btn-row { margin-top: 0.75rem; }

/* Keywords */
.kw-missing {
    background: #fef2f2; border: 1px solid #fecaca; color: #b91c1c;
    border-radius: 999px; padding: 0.28rem 0.7rem; font-size: 0.76rem; font-weight: 500;
    transition: all 0.15s;
}
.kw-missing:hover { background: #fee2e2; transform: scale(1.03); }
.kw-present {
    background: #E6F8F3; border: 1px solid rgba(21,195,154,0.3); color: #12A884;
    border-radius: 999px; padding: 0.28rem 0.7rem; font-size: 0.76rem; font-weight: 500;
    transition: all 0.15s;
}
.kw-present:hover { background: rgba(21,195,154,0.2); transform: scale(1.03); }

/* Alerts */
.warning-strip {
    background: #fffbeb; border-left: 3px solid #E8A735;
    border-radius: 0 10px 10px 0; padding: 0.65rem 0.9rem;
    font-size: 0.82rem; color: #78350f; margin-bottom: 0.5rem;
}
.ocr-strip {
    background: #eff6ff; border-left: 3px solid #3b82f6;
    border-radius: 0 10px 10px 0; padding: 0.65rem 0.9rem;
    font-size: 0.82rem; color: #1e40af; margin-bottom: 0.8rem;
}

/* Streamlit Overrides */
[data-testid="stFileUploader"] {
    border: 2px dashed #E8EAED !important; border-radius: 12px !important;
    background: #F7F8FA !important; transition: border-color 0.2s !important;
}
[data-testid="stFileUploader"]:hover { border-color: #15C39A !important; }

/* Buttons */
.stButton > button,
.stButton > button * {
    font-family: 'DM Sans', sans-serif !important;
    font-weight: 700 !important;
    font-size: 0.95rem !important;
    letter-spacing: 0.01em !important;
    text-rendering: optimizeLegibility !important;
    -webkit-font-smoothing: antialiased !important;
}
.stButton > button {
    background: linear-gradient(135deg, #8b5cf6, #7c3aed) !important;
    color: #f8f8f2 !important;
    border: 1px solid rgba(139, 92, 246, 0.35) !important;
    border-radius: 12px !important;
    white-space: nowrap !important;
    overflow: visible !important;
    text-overflow: clip !important;
    text-shadow: none !important;
    display: inline-block !important;
    box-shadow: 0 6px 18px rgba(124, 58, 237, 0.22) !important;
    transition: all 0.18s ease !important;
}
.stButton > button:hover {
    background: linear-gradient(135deg, #7c3aed, #6d28d9) !important;
    border-color: rgba(139, 92, 246, 0.55) !important;
    transform: translateY(-1px) !important;
    box-shadow: 0 8px 22px rgba(124, 58, 237, 0.28) !important;
}
.stButton > button:active { transform: translateY(0) !important; }

.stDownloadButton > button {
    border-radius: 10px !important;
    background: #0D0F11 !important;
    color: #FFFFFF !important; border: none !important;
    box-shadow: 0 2px 8px rgba(13,15,17,0.2) !important;
    font-weight: 700 !important; transition: all 0.2s ease !important;
}
.stDownloadButton > button:hover {
    box-shadow: 0 4px 16px rgba(13,15,17,0.3) !important;
    transform: translateY(-1px) !important;
}

/* Analytics & metric text */
[data-testid="stMetricLabel"],
[data-testid="stMetricValue"] {
    font-family: 'Instrument Serif', serif !important;
    font-weight: 400 !important;
    font-size: 48px !important;
    font-style: normal !important;
    letter-spacing: 0.02em !important;
    -webkit-font-smoothing: antialiased !important;
    text-rendering: optimizeLegibility !important;
}

/* Sidebar */
[data-testid="stSidebar"] { background-color: #1A1C1E; border-right: 1px solid #2A2D31; }
[data-testid="stSidebar"] p, [data-testid="stSidebar"] span, [data-testid="stSidebar"] label { color: #F7F8FA !important; }
[data-testid="stSidebar"] .card { background-color: #2D3136; border-color: #3C4148; }

textarea, input {
    border-radius: 10px !important; font-family: 'DM Sans', sans-serif !important;
    font-size: 0.84rem !important; border: 1px solid #E8EAED !important;
    transition: border-color 0.2s !important;
}
textarea:focus, input:focus { border-color: #15C39A !important; }
.stSpinner > div { border-top-color: #15C39A !important; }

/* Section completion cards */
.section-card {
    padding: 1rem; border-radius: 12px; text-align: center; color: white;
    font-weight: 700; transition: transform 0.2s; border: 1px solid transparent;
}
.section-card:hover { transform: scale(1.03); }
.section-card.present { background: #15C39A; border-color: #12A884; }
.section-card.missing { background: #E5534B; border-color: #C94A42; }

/* Radio Buttons for Navigation */
div[role="radiogroup"] {
    background: linear-gradient(135deg, #0a0a0f 0%, #111827 45%, #172554 100%);
    padding: 0.45rem 0.55rem;
    border-radius: 16px;
    border: 1px solid rgba(139,92,246,0.25);
    gap: 0.4rem;
    margin: 0 auto 1.25rem;
    width: min(100%, 1180px);
    display: flex;
    justify-content: center;
    align-items: center;
    box-shadow: 0 10px 24px rgba(10,10,15,0.18);
}
label[data-baseweb="radio"] {
    background: rgba(255,255,255,0.06);
    padding: 0.55rem 0.85rem;
    border-radius: 10px;
    transition: all 0.18s ease;
    flex: 1 1 auto;
    text-align: center;
    min-width: 110px;
    cursor: pointer;
    border: 1px solid transparent;
    display: flex;
    align-items: center;
    justify-content: center;
}
label[data-baseweb="radio"]:hover {
    background: rgba(255,255,255,0.12);
    border-color: rgba(20,184,166,0.18);
}
label[data-baseweb="radio"] div:first-child { display: none; }
label[data-baseweb="radio"] *,
label[data-baseweb="radio"] > div,
label[data-baseweb="radio"] span {
    color: #ffffff !important;
    opacity: 1 !important;
    font-family: 'Instrument Serif', serif !important;
    font-weight: 400 !important;
    font-size: 0.95rem !important;
    letter-spacing: 0.02em !important;
    white-space: nowrap !important;
    overflow: visible !important;
    text-overflow: clip !important;
    text-shadow: none !important;
    display: inline-block !important;
}
label[data-baseweb="radio"][aria-checked="true"] {
    background: linear-gradient(135deg, rgba(139,92,246,0.22), rgba(20,184,166,0.14));
    border-color: rgba(139,92,246,0.45);
    box-shadow: inset 0 0 0 1px rgba(255,255,255,0.08);
}
label[data-baseweb="radio"][aria-checked="true"] div:last-child {
    color: #ffffff !important;
    text-shadow: 0 1px 3px rgba(0,0,0,0.35);
}

/* ── Score Tooltip ── */
.score-tooltip-wrap {
    position: relative;
    cursor: pointer;
}
.score-tooltip-wrap .score-tooltip {
    visibility: hidden;
    opacity: 0;
    position: absolute;
    bottom: calc(100% + 10px);
    left: 50%;
    transform: translateX(-50%);
    background: #1A1C1E;
    color: #F7F8FA;
    border: 1px solid #3C4148;
    border-radius: 12px;
    padding: 0.75rem 1rem;
    font-size: 0.76rem;
    font-family: 'DM Sans', sans-serif;
    white-space: pre;
    line-height: 1.7;
    z-index: 999;
    box-shadow: 0 8px 24px rgba(0,0,0,0.35);
    pointer-events: none;
    transition: opacity 0.18s ease, visibility 0.18s ease;
    min-width: 200px;
}
.score-tooltip-wrap:hover .score-tooltip {
    visibility: visible;
    opacity: 1;
}
.score-tooltip-wrap .score-tooltip::after {
    content: '';
    position: absolute;
    top: 100%;
    left: 50%;
    transform: translateX(-50%);
    border: 6px solid transparent;
    border-top-color: #1A1C1E;
}

/* Model Selector Bar */
.model-bar {
    display: flex;
    align-items: center;
    gap: 1.5rem;
    padding: 0.5rem 3rem;
    margin: -0.5rem -3rem 0;
    background: #0a0a0f;
    border-bottom: 1px solid rgba(255,255,255,0.06);
}
.model-bar .stSelectbox label,
.model-bar .stToggle label {
    color: rgba(255,255,255,0.7) !important;
    font-size: 0.78rem !important;
}

/* Responsive */
@media (max-width: 900px) {
    .main .block-container { padding: 1rem; }
    .hero-wrap { border-radius: 0 0 18px 18px; padding: 2rem 1rem 1.25rem; margin: -1rem -1rem 1.25rem; }
    .hero-title { font-size: 2.4rem; }
    .hero-badge { display: inline-flex; margin-top: 0.9rem; }
    .rewrite-grid { grid-template-columns: 1fr; }
    .rewrite-pane.before { border-right: 0; border-bottom: 1px solid #E8EAED; }
    .pdf-frame, .pdf-shell { min-height: 560px; }
    .model-bar { padding: 0.5rem 1rem; margin: -0.5rem -1rem 0; }
}
"""

_SCORE_HISTORY_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "score_history.json")


def _load_score_history() -> list[dict]:
    try:
        if os.path.exists(_SCORE_HISTORY_PATH):
            with open(_SCORE_HISTORY_PATH, "r") as f:
                data = json.load(f)
            if isinstance(data, list):
                return data[-50:]  # cap at 50
    except Exception:
        pass
    return []


def _save_score_history(history: list[dict]) -> None:
    try:
        with open(_SCORE_HISTORY_PATH, "w") as f:
            json.dump(history[-50:], f)
    except Exception as e:
        print(f"Failed to save score history: {e}")


# providers
_PROVIDER_MODELS: dict[str, list[str]] = {
    "Gemini": ["gemma-4-31b-it", "gemini-2.5-flash", "gemini-2.5-pro"],
    "Claude": ["claude-4-5-sonnet-latest", "claude-4-5-haiku-latest"],
    "ChatGPT": ["gpt-4o-mini", "gpt-4o", "gpt-4.1-mini"],
    "Local": ["llama3"],
}

_DISPLAY_TO_PROVIDER: dict[str, str] = {
    "Gemini": "gemini",
    "Claude": "claude",
    "ChatGPT": "chatgpt",
    "Local": "local",
}


def _build_model_options() -> list[str]:
    options = []
    for provider, models in _PROVIDER_MODELS.items():
        for model in models:
            options.append(f"{provider} → {model}")
    return options


def _parse_model_selection(selection: str) -> tuple[str, str, str]:
    parts = selection.split(" → ", 1)
    display = parts[0].strip()
    model = parts[1].strip() if len(parts) > 1 else ""
    return display, _DISPLAY_TO_PROVIDER.get(display, "gemini"), model


@dataclass
class AppState:
    results:  dict[str, Any] | None = None
    accepted: dict[str, bool] = field(default_factory=dict)
    analysed: bool = False
    score_history: list[dict] = field(default_factory=list)
    generated_cv: str = ""
    generated_cover_letter: str = ""
    pdf_bytes: bytes | None = None
    last_file_hash: int | None = None
    trigger_sidebar: bool = False

def _init_session_state() -> None:
    default_state = AppState()
    for key, val in default_state.__dict__.items():
        if key not in st.session_state:
            st.session_state[key] = val
    # Load persistent score history from disk on first init
    if not st.session_state.score_history:
        st.session_state.score_history = _load_score_history()

def _get_score_config(score: int) -> dict[str, Any]:
    if score >= _SCORE_THRESHOLDS["strong"]["min"]:
        return _SCORE_THRESHOLDS["strong"]
    if score >= _SCORE_THRESHOLDS["medium"]["min"]:
        return _SCORE_THRESHOLDS["medium"]
    return _SCORE_THRESHOLDS["weak"]

def _render_hero() -> None:
    st.markdown("""
    <div class="hero-wrap">
        <div class="hero-eyebrow">Resume intelligence, reimagined</div>
        <h1 class="hero-title"><span class="title-prefix">RagsToRiches:</span><br><span class="accent">Smarter resumes, smarter opportunities.</span></h1>
        <p class="hero-sub">
            Review rewrite suggestions against the uploaded PDF, apply the changes you trust,
            then generate a CV from those decisions with or without a job description.
        </p>
    </div>
    """, unsafe_allow_html=True)

def _iter_rewrite_items(rewrites: dict[str, list[dict]]) -> list[tuple[str, int, dict]]:
    return [
        (section_name, i, item)
        for section_name, bullets in rewrites.items()
        for i, item in enumerate(bullets)
    ]

def _suggestion_key(section_name: str, i: int, item: dict) -> str:
    return item.get("id") or f"{section_name}_{i}"

def _actionable_items(rewrites: dict[str, list[dict]]) -> list[tuple[str, int, dict]]:
    return [
        (section_name, i, item)
        for section_name, i, item in _iter_rewrite_items(rewrites)
        if item.get("framework_used") not in ("none", "error")
        and item.get("original") != item.get("rewritten")
    ]

def _accepted_rewrite_map(rewrites: dict[str, list[dict]]) -> dict[str, str]:
    accepted_map: dict[str, str] = {}
    for section_name, i, item in _iter_rewrite_items(rewrites):
        if not isinstance(item, dict):
            continue
        key = _suggestion_key(section_name, i, item)
        if st.session_state.accepted.get(key) is True:
            accepted_map[item.get("original", "")] = item.get("rewritten", "")
    return accepted_map

def _generate_docx(markdown_text: str) -> bytes:
    try:
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        import re as _re
    except ImportError:
        raise RuntimeError("python-docx is not installed. Run: pip install python-docx")

    if not markdown_text:
        raise ValueError("Generated CV is empty.")
    if len(markdown_text.strip()) < 100:
        raise ValueError("Generated CV appears incomplete.")

    doc = Document()

    # Small margins to fit evreything in 2 pages
    for sec in doc.sections:
        sec.top_margin    = Inches(0.6)
        sec.bottom_margin = Inches(0.6)
        sec.left_margin   = Inches(0.75)
        sec.right_margin  = Inches(0.75)

    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)

    def _add_hr(doc):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after  = Pt(3)
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '4')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), 'CCCCCC')
        pBdr.append(bottom)
        pPr.append(pBdr)
        return p

    def _add_runs(p, raw_line: str, base_size: float = 10.5, base_bold: bool = False):
        import re as rr
        parts = rr.split(r'(\*\*.*?\*\*|\*.*?\*|_.*?_)', raw_line)
        for part in parts:
            if part.startswith('**') and part.endswith('**') and len(part) > 4:
                run = p.add_run(part[2:-2])
                run.bold = True
            elif (part.startswith('*') and part.endswith('*') and len(part) > 2) or \
                 (part.startswith('_') and part.endswith('_') and len(part) > 2):
                run = p.add_run(part[1:-1])
                run.italic = True
            else:
                run = p.add_run(part)
                run.bold = base_bold
            run.font.size = Pt(base_size)

    lines = markdown_text.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line.strip():
            i += 1
            continue

        # Candidate name
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after  = Pt(2)
            run = p.add_run(line[2:].strip())
            run.bold = True
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(0x0D, 0x0F, 0x11)
            i += 1
            continue

        # Section headings 
        if line.startswith('## '):
            heading_text = line[3:].strip().upper()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after  = Pt(0)
            run = p.add_run(heading_text)
            run.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x0D, 0x0F, 0x11)
            _add_hr(doc)
            i += 1
            continue

        # Role titles 
        if line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after  = Pt(1)
            _add_runs(p, line[4:].strip(), base_size=10.5, base_bold=True)
            i += 1
            continue

        # Thin divider
        if _re.match(r'^-{3,}$|^\*{3,}$|^_{3,}$', line.strip()):
            _add_hr(doc)
            i += 1
            continue

        # Bullet point
        bullet_match = _re.match(r'^(\s*)[-*•]\s+(.*)', line)
        if bullet_match:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after  = Pt(2)
            p.paragraph_format.left_indent  = Inches(0.2)
            _add_runs(p, bullet_match.group(2).strip(), base_size=10.5)
            i += 1
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after  = Pt(3)
        _add_runs(p, line.strip(), base_size=10.5)
        i += 1

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()

def _generate_pdf(markdown_text: str) -> bytes | None:
    import tempfile
    import os

    try:
        from markdown_pdf import MarkdownPdf
        from markdown_pdf import Section
    except ImportError:
        try:
            from markdown_pdf import MarkdownPdf
            from markdown_pdf.Section import Section
        except ImportError:
            return None

    pdf = MarkdownPdf(toc_level=0)
    pdf.add_section(Section(markdown_text))

    # markdown-pdf .save() 
    tmp_path = os.path.join(tempfile.gettempdir(), "ragstoriches_cv.pdf")
    pdf.save(tmp_path)

    with open(tmp_path, "rb") as f:
        data = f.read()

    try:
        os.remove(tmp_path)
    except OSError:
        pass

    return data

def _severity_color(severity: str, active: bool = False) -> tuple[float, float, float]:
    if active:
        return (0.04, 0.45, 0.28)
    if severity == "red":
        return (0.86, 0.24, 0.18)
    if severity == "green":
        return (0.10, 0.58, 0.33)
    return (0.92, 0.62, 0.05)

def _search_phrases(text: str) -> list[str]:
    cleaned = " ".join(text.split())
    words = cleaned.split()
    phrases = [cleaned]
    if len(words) > 18:
        phrases.append(" ".join(words[:18]))
    if len(words) > 10:
        phrases.append(" ".join(words[:10]))
    if len(words) > 10:
        phrases.append(" ".join(words[-10:]))
    return [phrase for phrase in phrases if len(phrase) >= 20]

@st.cache_data(show_spinner=False)
def _highlight_pdf_bytes(pdf_bytes: bytes, rewrite_items: tuple[tuple[str, str, str, str, str], ...]) -> bytes:
    try:
        import fitz
    except Exception:
        return pdf_bytes

    try:
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    except Exception:
        return pdf_bytes

    for item_id, text, severity, reasoning, rewritten in rewrite_items:
        found = False
        for phrase in _search_phrases(text):
            if found:
                break
            for page in doc:
                matches = page.search_for(phrase, quads=True)
                if not matches:
                    continue
                color = _severity_color(severity, active=False)
                for quad in matches[:3]:
                    annot = page.add_highlight_annot(quad)
                    annot.set_colors(stroke=color)
                    # popup: severity + suggestion + reasoning
                    if rewritten and rewritten != text:
                        popup_text = (
                            f"[{severity.upper()}] Suggested rewrite:\n{rewritten}\n\n"
                            f"Why: {reasoning}"
                        )
                    elif reasoning:
                        popup_text = f"[{severity.upper()}] {reasoning}"
                    else:
                        popup_text = f"Rewrite suggestion: {item_id}"
                    annot.set_info(content=popup_text, title="RAGsToRiches")
                    annot.update(opacity=0.32)
                found = True
                break

    output = io.BytesIO()
    doc.save(output, garbage=4, deflate=True)
    doc.close()
    return output.getvalue()

def _render_pdf_viewer(pdf_bytes: bytes | None, rewrites: dict[str, list[dict]]) -> None:
    if not pdf_bytes:
        st.markdown('<div class="card muted">Upload and analyse a PDF to see highlighted rewrite targets.</div>', unsafe_allow_html=True)
        return

    highlight_payload = tuple(
        (
            _suggestion_key(section_name, i, item),
            item.get("highlight_text") or item.get("original", ""),
            item.get("severity", "yellow"),
            item.get("reasoning", ""),
            item.get("rewritten", ""),
        )
        for section_name, i, item in _actionable_items(rewrites)
    )
    rendered = _highlight_pdf_bytes(pdf_bytes, highlight_payload) if highlight_payload else pdf_bytes
    encoded = base64.b64encode(rendered).decode("utf-8")
    st.markdown(
        f"""
        <div class="pdf-shell">
            <iframe class="pdf-frame" src="data:application/pdf;base64,{encoded}#toolbar=1&navpanes=0"></iframe>
        </div>
        """,
        unsafe_allow_html=True,
    )

def _render_rewrites_tab(rewrites: dict[str, list[dict]], pdf_bytes: bytes | None) -> None:
    actionable = _actionable_items(rewrites)
    if not rewrites:
        st.markdown('<p class="muted">No experience or project sections were found.</p>', unsafe_allow_html=True)
        return

    accepted_count = sum(
        1 for section_name, i, item in actionable
        if st.session_state.accepted.get(_suggestion_key(section_name, i, item)) is True
    )
    dismissed_count = sum(
        1 for section_name, i, item in actionable
        if st.session_state.accepted.get(_suggestion_key(section_name, i, item)) is False
    )

    st.markdown(
        f"""
        <div class="review-toolbar" style="margin-bottom:0.75rem">
            <span class="status-chip">{len(actionable)} suggested changes</span>
            <span class="status-chip">{accepted_count} accepted</span>
            <span class="status-chip">{dismissed_count} dismissed</span>
        </div>
        """,
        unsafe_allow_html=True,
    )

    tool_a, tool_b, _ = st.columns([1, 1, 4])
    with tool_a:
        if st.button("Accept all", use_container_width=True):
            for section_name, i, item in actionable:
                st.session_state.accepted[_suggestion_key(section_name, i, item)] = True
            st.session_state.generated_cv = ""
            st.rerun()
    with tool_b:
        if st.button("Clear decisions", use_container_width=True):
            st.session_state.accepted = {}
            st.session_state.generated_cv = ""
            st.rerun()

    pdf_col, review_col = st.columns([1.18, 1], gap="large")
    with pdf_col:
        st.markdown('<span class="section-label">Highlighted PDF</span>', unsafe_allow_html=True)
        _render_pdf_viewer(pdf_bytes, rewrites)

    with review_col:
        st.markdown('<span class="section-label">Rewrite Decisions</span>', unsafe_allow_html=True)
        st.markdown('<div class="suggestion-scroll">', unsafe_allow_html=True)

        if not actionable:
            st.markdown(
                '<div class="card muted">No rewrite-worthy sentences were detected. Header/date lines were skipped.</div>',
                unsafe_allow_html=True,
            )

        current_section = ""
        for section_name, i, item in actionable:
            if section_name != current_section:
                current_section = section_name
                st.markdown(f'<div class="result-section-head">{html.escape(section_name.title())}</div>', unsafe_allow_html=True)

            key = _suggestion_key(section_name, i, item)
            state = st.session_state.accepted.get(key)
            state_class = "accepted" if state is True else "dismissed" if state is False else ""
            state_text = "Accepted" if state is True else "Dismissed" if state is False else "Needs decision"
            safe_original = html.escape(item.get("original", ""))
            safe_rewritten = html.escape(item.get("rewritten", ""))
            safe_reasoning = html.escape(item.get("reasoning", ""))
            fw_badge = html.escape(item.get("framework_used", ""))

            severity = item.get("severity", "yellow")
            sev_dot = f'<span class="severity-dot {severity}"></span>'

            st.markdown(
                f"""
                <div class="suggestion-card {state_class}">
                    <div class="suggestion-head">
                        <span class="suggestion-title">{sev_dot}{html.escape(state_text)}</span>
                        <span class="fw-badge">{fw_badge}</span>
                    </div>
                    <div class="rewrite-grid">
                        <div class="rewrite-pane before">
                            <span class="pane-label">Original</span>
                            <p class="rewrite-text">{safe_original}</p>
                        </div>
                        <div class="rewrite-pane after">
                            <span class="pane-label">Suggested rewrite</span>
                            <p class="rewrite-text">{safe_rewritten}</p>
                        </div>
                    </div>
                    <div class="reasoning-row"> {safe_reasoning}</div>
                """,
                unsafe_allow_html=True,
            )

            if state is True:
                st.markdown('<div class="decision-bar accepted">This rewrite will be used in CV generation.</div></div>', unsafe_allow_html=True)
            elif state is False:
                st.markdown('<div class="decision-bar dismissed">This rewrite will be ignored in CV generation.</div></div>', unsafe_allow_html=True)
            else:
                st.markdown('</div>', unsafe_allow_html=True)

            st.markdown('<div class="suggestion-btn-row">', unsafe_allow_html=True)
            btn_a, btn_b = st.columns(2)
            with btn_a:
                if st.button("Accept", key=f"acc_{key}", use_container_width=True):
                    st.session_state.accepted[key] = True
                    st.session_state.generated_cv = ""
                    st.rerun()
            with btn_b:
                if st.button("Dismiss", key=f"rej_{key}", use_container_width=True):
                    st.session_state.accepted[key] = False
                    st.session_state.generated_cv = ""
                    st.rerun()
            st.markdown('</div>', unsafe_allow_html=True)

        st.markdown('</div>', unsafe_allow_html=True)

def _render_keywords_tab(jd_kws: list[str], missing: list[str], present: list[str], freqs: dict[str, int]) -> None:
    if not jd_kws:
        st.markdown(
            '<p style="color:#a1a1aa;padding:1rem 0">Paste a job description above and re-analyse to see keyword gaps.</p>',
            unsafe_allow_html=True,
        )
        return

    coverage = int(len(present) / len(jd_kws) * 100) if jd_kws else 0
    st.markdown(f"""
    <div class="card" style="margin-bottom:1.5rem">
        <span class="section-label">Coverage</span>
        <p style="font-size:2.2rem;font-weight:400;color:#0D0F11;margin:0;font-family:'Instrument Serif',serif;letter-spacing:0.02em;text-rendering:optimizeLegibility;-webkit-font-smoothing:antialiased;">
            {coverage}<span style="font-size:1rem;color:#6B7280;font-weight:400;font-family:'Instrument Serif',serif;">%</span>
        </p>
        <p style="font-size:0.82rem;color:#6B7280;margin:2px 0 0">
            {len(present)} of {len(jd_kws)} JD keywords present in your resume
        </p>
    </div>
    """, unsafe_allow_html=True)

    miss_col, have_col = st.columns(2, gap="large")

    with miss_col:
        st.markdown(f'<span class="section-label"> Missing ({len(missing)})</span>', unsafe_allow_html=True)
        chips = "".join(f'<span class="kw-missing">{html.escape(kw)}</span>' for kw in missing)
        st.markdown(
            f'<div class="kw-wrap">{chips or "<i style=\'color:#a1a1aa\'>None — great coverage!</i>"}</div>',
            unsafe_allow_html=True,
        )

    with have_col:
        st.markdown(f'<span class="section-label"> Present ({len(present)})</span>', unsafe_allow_html=True)
        chips = "".join(f'<span class="kw-present">{html.escape(kw)} <span style="opacity:0.7;font-size:0.8em">({freqs.get(kw, 0)})</span></span>' for kw in present)
        st.markdown(
            f'<div class="kw-wrap">{chips or "<i style=\'color:#a1a1aa\'>No matches found</i>"}</div>',
            unsafe_allow_html=True,
        )

def main() -> None:
    if not BACKEND_AVAILABLE:
        st.error("Missing backend logic. Ensure parser.py, analyser.py and router.py exist.")
        st.stop()

    st.set_page_config(**_PAGE_CONFIG)
    _init_session_state()
    if _CUSTOM_CSS.strip():
        st.markdown(f"<style>{_CUSTOM_CSS}</style>", unsafe_allow_html=True)

    _render_hero()

    #Model selector bar
    model_options = _build_model_options()
    bar_left, bar_right = st.columns([1, 1], gap="large")
    with bar_left:
        model_selection = st.selectbox(
            "LLM Model",
            options=model_options,
            index=0,
            help="Select provider and model. Ensure the corresponding API key is in your .env file.",
            label_visibility="collapsed",
        )
    with bar_right:
        use_critic = st.toggle(
            "Agentic Self-Correction",
            value=False,
            help="Runs a self-correction loop on rewritten bullets. Slower but higher quality results.",
        )

    display_provider, selected_provider, selected_model = _parse_model_selection(model_selection)

    local_endpoint = "http://localhost:11434/api/chat"
    if selected_provider == "local":
        local_endpoint = st.text_input(
            "Local API Endpoint",
            value="http://localhost:11434/api/chat",
            help="Default Ollama endpoint.",
        )
    else:
        key_missing = False
        if selected_provider == "gemini" and not os.environ.get("GEMINI_API_KEY"):
            key_missing = True
        elif selected_provider == "claude" and not os.environ.get("ANTHROPIC_API_KEY"):
            key_missing = True
        elif selected_provider == "chatgpt" and not os.environ.get("OPENAI_API_KEY"):
            key_missing = True
        if key_missing:
            st.error(f"Missing API key for {display_provider}. Please configure it in your .env file.")
            st.stop()

    st.markdown('<div class="setup-band">', unsafe_allow_html=True)
    st.markdown('<hr class="slim-divider">', unsafe_allow_html=True)
    left, right = st.columns([0.95, 1.25], gap="large")

    with left:
        st.markdown('<span class="section-label">Resume PDF</span>', unsafe_allow_html=True)
        uploaded_file = st.file_uploader("Upload Resume", type="pdf", label_visibility="collapsed", key="pdf_upload")
        if uploaded_file:
            st.success(f"{uploaded_file.name}")

    with right:
        st.markdown('<span class="section-label">Job Description <span style="color:#6B7280;font-weight:500;text-transform:none;letter-spacing:0">(optional for ATS matching)</span></span>', unsafe_allow_html=True)
        job_description = st.text_area(
            "Job Description",
            placeholder="Paste a full job description for keyword matching, or leave blank to improve the CV from rewrite decisions only.",
            height=152,
            label_visibility="collapsed",
        )
    st.markdown('</div>', unsafe_allow_html=True)

    if uploaded_file:
        uploaded_bytes = uploaded_file.getvalue()
        st.session_state.pdf_bytes = uploaded_bytes
        analyse_clicked = st.button("Analyse resume ✨", use_container_width=True)
        
        current_hash = hash(uploaded_bytes)
        if st.session_state.get("last_file_hash") != current_hash:
             st.session_state.results = None
             st.session_state.analysed = False
             st.session_state.last_file_hash = current_hash
             st.session_state.accepted = {}
             st.session_state.generated_cv = ""
             st.session_state.generated_cover_letter = ""

        if analyse_clicked:
            with st.spinner("Parsing resume..."):
                parsed = extract_text(uploaded_bytes)

            if not parsed.sections and parsed.warnings:
                for w in parsed.warnings:
                    st.error(w)
                st.stop()

            with st.spinner(f"Running AI analysis via {display_provider} ({selected_model}) — this may take a few minutes..."):
                try:
                    results = analyse_resume(
                        resume=parsed, 
                        job_description=job_description,
                        provider=selected_provider,
                        local_endpoint=local_endpoint,
                        use_critic=use_critic,
                        model=selected_model,
                    )
                    results["parsed_resume_obj"] = parsed
                    st.session_state.parsed_resume = parsed
                    st.session_state.pdf_bytes = uploaded_bytes
                    st.session_state.results  = results
                    # Store full breakdown dict in history
                    score_data = results["score"]
                    st.session_state.score_history.append(score_data)
                    if len(st.session_state.score_history) > 50:
                        st.session_state.score_history = st.session_state.score_history[-50:]
                    _save_score_history(st.session_state.score_history)
                    st.session_state.accepted = {}
                    st.session_state.generated_cv = ""
                    st.session_state.generated_cover_letter = ""
                    st.session_state.analysed = True
                    st.session_state.trigger_sidebar = True
                except Exception as e:
                    st.error(f"Analysis failed. Did you configure {display_provider} correctly? Error: {e}")
                    st.stop()

            st.rerun()
    else:
        st.markdown(
            '<p class="muted" style="text-align:center;padding:0.5rem 0">Upload a PDF resume above to get started.</p>',
            unsafe_allow_html=True,
        )

    if st.session_state.results:
        r = st.session_state.results

        st.markdown("<hr class='slim-divider'>", unsafe_allow_html=True)

        if r.get("ocr_used"):
            st.markdown(
                '<div class="ocr-strip">🔍 No text layer detected — OCR was used. '
                'Accuracy may be slightly lower on styled or image-heavy PDFs.</div>',
                unsafe_allow_html=True,
            )

        score_data = r.get("score", {"total": 0, "base": 30, "sections": 0, "keywords": 0, "bullet_quality": 0, "action_verbs": 0, "warnings": 0})
        if isinstance(score_data, int):
            score = score_data
            tooltip_html = "Score breakdown not available"
        else:
            score = score_data.get("total", 0)
            tooltip_lines = [
                f"Base:            {score_data.get('base', 0)}",
                f"Sections:      +{score_data.get('sections', 0)}",
                f"Keywords:      +{score_data.get('keywords', 0)}",
                f"Bullet Quality: +{score_data.get('bullet_quality', 0)}",
                f"Action Verbs:  +{score_data.get('action_verbs', 0)}",
                f"Warnings:       {score_data.get('warnings', 0)}",
                f"{'─' * 26}",
                f"Total:          {score}/100",
            ]
            tooltip_html = html.escape("\n".join(tooltip_lines))

        score_cfg = _get_score_config(score)

        st.markdown(f"""
        <div class="card" style="margin-bottom: 1rem; background: #FFFFFF; border: 1px solid #E8EAED; border-radius: 16px; padding: 1rem 1.1rem;">
            <span class="section-label">Resume Score</span>
            <div class="score-tooltip-wrap">
                <div class="score-ring-wrap" style="gap: 1rem; align-items: center;">
                    <div class="score-number" style="color:{score_cfg['color']};">{score}</div>
                    <div class="score-meta">
                        <span class="score-label-text">{score_cfg['label']}</span>
                        <span class="score-sub">out of 100 · hover for breakdown</span>
                    </div>
                </div>
                <div class="score-tooltip">{tooltip_html}</div>
            </div>
        </div>
        """, unsafe_allow_html=True)

        with st.sidebar:
            st.markdown(f"""
            <div class="card" style="margin-top:2rem">
                <span class="section-label">Resume Score</span>
                <div class="score-tooltip-wrap">
                    <div class="score-ring-wrap">
                        <div class="score-number" style="color:{score_cfg['color']}">{score}</div>
                        <div class="score-meta">
                            <span class="score-label-text">{score_cfg['label']}</span>
                            <span class="score-sub">out of 100</span>
                        </div>
                    </div>
                    <div class="score-tooltip">{tooltip_html}</div>
                </div>
            </div>
            """, unsafe_allow_html=True)
            
            contact = r.get("contact", {})
            if contact:
                chips = "".join(
                    f'<span class="contact-chip"><b>{html.escape(k.title())}</b>{html.escape(str(v))}</span>'
                    for k, v in contact.items()
                )
                st.markdown(f"""
                <div class="card" style="margin-top:1rem">
                    <span class="section-label">Contact Detected</span>
                    <div class="contact-grid">{chips}</div>
                </div>
                """, unsafe_allow_html=True)

        warnings = [w for w in r.get("warnings", []) if "not detected" in w.lower() or "corrupt" in w.lower()]
        if warnings:
            st.markdown('<span class="section-label" style="margin-top:0.5rem">Parser Notes</span>', unsafe_allow_html=True)
            for w in warnings:
                st.markdown(f'<div class="warning-strip">{html.escape(w)}</div>', unsafe_allow_html=True)

        with st.sidebar:
            st.markdown('<hr style="border-top:1px solid #3C4148">', unsafe_allow_html=True)
            st.markdown('<span class="section-label" style="color:#A1A5AB;">Parser Debug</span>', unsafe_allow_html=True)
            
            parsed_sections = r.get("sections", {})
            for sec in ["EXPERIENCE", "EDUCATION", "SKILLS", "PROJECTS"]:
                icon = "✓" if sec in parsed_sections else "✗"
                color = "#15C39A" if sec in parsed_sections else "#E5534B"
                st.markdown(f"<span style='color:{color}'>{icon} {sec.title()}</span>", unsafe_allow_html=True)
                
            with st.expander("Raw Parsed Output", expanded=False):
                st.json(parsed_sections)

        nav_options = ["Rewrite Suggestions", "Keyword Gap", "Extracted Sections", "Tailored CV", "Cover Letter", "Analytics"]
        selected_view = st.radio("", nav_options, horizontal=True, label_visibility="collapsed")

        if selected_view == "Rewrite Suggestions":
            _render_rewrites_tab(r.get("rewrites", {}), st.session_state.get("pdf_bytes"))

        elif selected_view == "Keyword Gap":
            jd_kws  = r.get("jd_keywords", [])
            missing = r.get("missing_keywords", [])
            present = [kw for kw in jd_kws if kw not in missing]
            freqs = r.get("keyword_frequencies", {})
            _render_keywords_tab(jd_kws, missing, present, freqs)

        elif selected_view == "Extracted Sections":
            sections = r.get("sections", {})
            for section_name, lines in sections.items():
                with st.expander(section_name.title(), expanded=(section_name == "EXPERIENCE")):
                    st.code("\n".join(lines), language=None)

        elif selected_view == "Tailored CV":
            st.markdown('<span class="section-label">Generate Tailored CV</span>', unsafe_allow_html=True)
            st.markdown(
                '<p class="muted">The generated CV applies accepted rewrites, ignores dismissed rewrites, and uses the job description only when one was provided during analysis.</p>',
                unsafe_allow_html=True,
            )
            if st.button("Generate CV", key="gen_cv_btn"):
                acc_rewrites_map = _accepted_rewrite_map(r.get("rewrites", {}))

                with st.spinner("Generating CV..."):
                    cv_text = generate_cv(
                        st.session_state.parsed_resume,
                        job_description,
                        acc_rewrites_map,
                        selected_provider,
                        local_endpoint,
                        rewrite_suggestions=r.get("rewrites", {}),
                        rewrite_decisions=st.session_state.accepted,
                        model=selected_model,
                    )
                    st.session_state.generated_cv = cv_text

            if st.session_state.get("generated_cv"):
                st.markdown("### ✏️ Edit Your CV")
                st.markdown(
                    '<p class="muted">Make any final edits below. Your changes will be reflected in the downloaded files.</p>',
                    unsafe_allow_html=True,
                )
                edited_cv = st.text_area(
                    "Edit CV",
                    value=st.session_state.generated_cv,
                    height=450,
                    label_visibility="collapsed",
                    key="cv_editor",
                )
                # Sync edits back to session state
                if edited_cv != st.session_state.generated_cv:
                    st.session_state.generated_cv = edited_cv

                st.markdown("### Preview")
                st.markdown(st.session_state.generated_cv)

                dl_col1, dl_col2, dl_col3 = st.columns(3)
                with dl_col1:
                    st.download_button(
                        "📄 Download Markdown",
                        st.session_state.generated_cv,
                        "tailored_cv.md",
                        key="dl_md",
                    )

                with dl_col2:
                    try:
                        docx_bytes = _generate_docx(st.session_state.generated_cv)
                        st.download_button(
                            "📝 Download DOCX",
                            docx_bytes,
                            "tailored_cv.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            key="dl_docx",
                        )
                    except Exception as e:
                        st.error(f"DOCX export failed: {e}")

                with dl_col3:
                    try:
                        pdf_dl_bytes = _generate_pdf(st.session_state.generated_cv)
                        if pdf_dl_bytes:
                            st.download_button(
                                "📕 Download PDF",
                                pdf_dl_bytes,
                                "tailored_cv.pdf",
                                mime="application/pdf",
                                key="dl_pdf",
                            )
                    except Exception as e:
                        st.error(f"PDF export failed: {e}")

        elif selected_view == "Cover Letter":
            st.markdown('<span class="section-label">Generate Cover Letter</span>', unsafe_allow_html=True)
            st.markdown(
                '<p class="muted">Generate a professional cover letter tailored to the job description. '
                'A job description is recommended for best results but not required.</p>',
                unsafe_allow_html=True,
            )
            if st.button("Generate Cover Letter", key="gen_cl_btn"):
                with st.spinner("Generating cover letter..."):
                    cl_text = generate_cover_letter(
                        st.session_state.parsed_resume,
                        job_description,
                        selected_provider,
                        local_endpoint,
                        model=selected_model,
                    )
                    st.session_state.generated_cover_letter = cl_text

            if st.session_state.get("generated_cover_letter"):
                st.markdown("### ✏️ Edit Your Cover Letter")
                st.markdown(
                    '<p class="muted">Make any final edits below. Your changes will be reflected in the downloaded files.</p>',
                    unsafe_allow_html=True,
                )
                edited_cl = st.text_area(
                    "Edit Cover Letter",
                    value=st.session_state.generated_cover_letter,
                    height=400,
                    label_visibility="collapsed",
                    key="cl_editor",
                )
                if edited_cl != st.session_state.generated_cover_letter:
                    st.session_state.generated_cover_letter = edited_cl

                st.markdown("### Preview")
                st.markdown(st.session_state.generated_cover_letter)

                cl_col1, cl_col2, cl_col3 = st.columns(3)
                with cl_col1:
                    st.download_button(
                        "📄 Download Markdown",
                        st.session_state.generated_cover_letter,
                        "cover_letter.md",
                        key="dl_cl_md",
                    )

                with cl_col2:
                    try:
                        cl_docx = _generate_docx(st.session_state.generated_cover_letter)
                        st.download_button(
                            "📝 Download DOCX",
                            cl_docx,
                            "cover_letter.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            key="dl_cl_docx",
                        )
                    except Exception as e:
                        st.error(f"DOCX export failed: {e}")

                with cl_col3:
                    try:
                        cl_pdf = _generate_pdf(st.session_state.generated_cover_letter)
                        if cl_pdf:
                            st.download_button(
                                "📕 Download PDF",
                                cl_pdf,
                                "cover_letter.pdf",
                                mime="application/pdf",
                                key="dl_cl_pdf",
                            )
                    except Exception as e:
                        st.error(f"PDF export failed: {e}")

        elif selected_view == "Analytics":
            st.markdown('<span class="section-label">User Analytics Dashboard</span>', unsafe_allow_html=True)
            if not r:
                st.info("Run your first analysis to see progress here.")
            else:
                st.markdown("### Resume Score History")
                history = st.session_state.score_history
                if len(history) > 1:
                    try:
                        import altair as alt
                        import pandas as pd

                        # Extract totals
                        chart_data = []
                        for idx, entry in enumerate(history):
                            if isinstance(entry, dict):
                                chart_data.append({"Attempt": idx + 1, "Score": entry.get("total", 0)})
                            else:
                                chart_data.append({"Attempt": idx + 1, "Score": int(entry)})

                        df = pd.DataFrame(chart_data)

                        chart = (
                            alt.Chart(df)
                            .mark_line(
                                point=alt.OverlayMarkDef(filled=True, size=60, color="#8b5cf6"),
                                strokeWidth=2.5,
                                color="#8b5cf6",
                            )
                            .encode(
                                x=alt.X(
                                    "Attempt:Q",
                                    title="Attempt",
                                    scale=alt.Scale(domain=[1, max(len(history), 2)]),
                                    axis=alt.Axis(tickMinStep=1, format="d"),
                                ),
                                y=alt.Y(
                                    "Score:Q",
                                    title="Resume Score",
                                    scale=alt.Scale(domain=[0, 100]),
                                ),
                                tooltip=["Attempt:Q", "Score:Q"],
                            )
                            .properties(height=200)
                            .configure_axis(
                                labelFontSize=11,
                                titleFontSize=12,
                            )
                        )
                        st.altair_chart(chart, use_container_width=True)
                    except ImportError:
                        # Fallback if altair not installed
                        scores = [e.get("total", 0) if isinstance(e, dict) else int(e) for e in history]
                        st.line_chart(scores, height=150)
                elif len(history) == 1:
                    entry = history[0]
                    single_score = entry.get("total", 0) if isinstance(entry, dict) else int(entry)
                    st.info(f"First analysis score: **{single_score}/100**. Run more analyses to see a trend line.")
                else:
                    st.info("Run multiple analyses to see a trend line.")

                st.markdown("### Per-Run Metrics")
                
                jd_kws = r.get("jd_keywords", [])
                missing = r.get("missing_keywords", [])
                present = [kw for kw in jd_kws if kw not in missing]
                coverage = (len(present) / len(jd_kws) * 100) if jd_kws else 0
                
                rewrites = r.get("rewrites", {})
                improved_bullets = sum(
                    1 for section in rewrites.values()
                    for item in section
                    if isinstance(item, dict)
                    and item.get("framework_used") not in ("none", "error")
                    and item.get("original") != item.get("rewritten")
                )
                
                cols = st.columns(2)
                with cols[0]:
                    if r.get("no_jd_provided"):
                        st.metric("Keyword Coverage", "N/A", help="No Job Description provided.")
                    else:
                        st.metric("Keyword Coverage", f"{coverage:.0f}%")
                with cols[1]:
                    st.metric("Bullets Improved", improved_bullets)

                st.markdown("### Section Completion")
                st.markdown("This highlights which sections are present in your resume.")
                
                cols = st.columns(3)
                sections = ["EXPERIENCE", "PROJECTS", "EDUCATION"]
                for i, sec in enumerate(sections):
                    with cols[i]:
                        has_sec = sec in r.get("sections", {})
                        css_class = "present" if has_sec else "missing"
                        status = "✓ Present" if has_sec else "✗ Missing"
                        st.markdown(f"""
                        <div class="section-card {css_class}">
                            {sec}<br>
                            <span style="font-size: 0.8em;">{status}</span>
                        </div>
                        """, unsafe_allow_html=True)

    if st.session_state.get("trigger_sidebar"):
        import streamlit.components.v1 as components
        components.html(
            """
            <script>
                const expandBtn = window.parent.document.querySelector('[data-testid="collapsedControl"]');
                if (expandBtn) {
                    expandBtn.click();
                }
            </script>
            """,
            height=0,
            width=0,
        )
        st.session_state.trigger_sidebar = False

if __name__ == "__main__":
    main()