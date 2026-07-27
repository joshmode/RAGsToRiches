import io
import os
import re
import tempfile


def generate_docx(md_text: str) -> bytes:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    if not md_text or len(md_text.strip()) < 20:
        raise ValueError("Generated document is empty.")

    doc = Document()

    # tight margins to keep resume under 2 page
    for sec in doc.sections:
        sec.top_margin    = Inches(0.6)
        sec.bottom_margin = Inches(0.6)
        sec.left_margin   = Inches(0.75)
        sec.right_margin  = Inches(0.75)

    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)

    def _add_hr():
        """add a thin hr via paragraph border."""
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after  = Pt(3)
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '4')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), 'CCCCCC')
        pBdr.append(bottom)
        pPr.append(pBdr)

    def _add_runs(p, raw_line: str, base_size: float = 10.5, base_bold: bool = False):
        """parse inline **bold** and *italic* markers."""
        for part in re.split(r'(\*\*.*?\*\*|\*.*?\*|_.*?_)', raw_line):
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
                run.italic = False
            elif (part.startswith('*') and part.endswith('*')) or (part.startswith('_') and part.endswith('_')):
                run = p.add_run(part[1:-1])
                run.bold = base_bold
                run.italic = True
            else:
                run = p.add_run(part)
                run.bold = base_bold
                run.italic = False
            run.font.size = Pt(base_size)

    for line in md_text.split('\n'):
        line = line.rstrip()
        if not line.strip():
            continue

        # h1: candidate name
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line[2:].strip())
            run.bold = True
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(0x0D, 0x0F, 0x11)

        # h2: section headings
        elif line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            run = p.add_run(line[3:].strip().upper())
            run.bold = True
            run.font.size = Pt(10)
            _add_hr()

        # h3: role/project titles
        elif line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            _add_runs(p, line[4:].strip(), base_bold=True)

        elif re.match(r'^-{3,}$|^\*{3,}$|^_{3,}$', line.strip()):
            _add_hr()

        else:
            m = re.match(r'^(\s*)[-*•]\s+(.*)', line)
            p = doc.add_paragraph(style='List Bullet') if m else doc.add_paragraph()
            if m:
                p.paragraph_format.left_indent = Inches(0.2)
                _add_runs(p, m.group(2).strip())
            else:
                _add_runs(p, line.strip())

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_pdf(md_text: str) -> bytes:
    from markdown_pdf import MarkdownPdf
    try:
        from markdown_pdf import Section
    except ImportError:
        from markdown_pdf.Section import Section

    pdf = MarkdownPdf(toc_level=0)
    pdf.add_section(Section(md_text))
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
        path = tmp.name
    try:
        pdf.save(path)
        with open(path, "rb") as f:
            return f.read()
    finally:
        if os.path.exists(path):
            os.remove(path)
